import csv
import hashlib
import json
import locale
import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from applications.extensions import db
from applications.models.analysis import Analysis
from applications.models.jiangxi_seed import JiangxiMinePlot
from applications.models.project import (
    Project,
    ProjectActivityLog,
    ProjectBackupRecord,
    ProjectDataset,
    ProjectExportRecord,
    ProjectMineBinding,
)
from applications.project_hub.tbbh_identity import normalize_tbbh

DEFAULT_JIANGXI_PROJECT_NAME = "江西矿山生态修复监测项目"
DEFAULT_JIANGXI_REGION = "江西省"
DEFAULT_JIANGXI_REMARK = "system_seed:jiangxi_tbbh"
DEFAULT_MONITOR_START_YEAR = 2017
DEFAULT_MONITOR_END_YEAR = 2025
REQUIRED_COLUMNS = ("主体编号", "地市", "区县", "矿山位置", "Area")
WORKBOOK_REQUIRED_COLUMNS = ("FID", "TBBH", "SHI", "XIAN", "面积", "矿山位置_")
WORKBOOK_MONITOR_START_YEAR = 2013
WORKBOOK_MONITOR_END_YEAR = 2025
PLOT_CONTENT_FIELDS = (
    "tbbh",
    "city",
    "county",
    "location_text",
    "plot_code",
    "restored_plot_code",
    "plot_category",
    "repair_status",
    "repair_mode",
    "completed_at",
    "area",
    "untreated_area",
    "center_lng",
    "center_lat",
    "closed_year",
    "plot_attr",
)


def _plot_content_signature(item):
    if isinstance(item, dict):
        return tuple(item.get(field) for field in PLOT_CONTENT_FIELDS)
    return tuple(getattr(item, field) for field in PLOT_CONTENT_FIELDS)


def _normalize_text(value):
    text = str(value or "").strip()
    return text or None


def _extract_year(value):
    match = re.search(r"(19|20)\d{2}", str(value or "").strip())
    return int(match.group(0)) if match else None


def _parse_float(value, *, field_name):
    text = str(value or "").strip()
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{field_name} 不是合法数值: {text}") from exc


def _candidate_encodings():
    preferred = locale.getpreferredencoding(False) or "gbk"
    return [preferred, "gbk", "gb18030", "utf-8-sig", "utf-8"]


def _read_csv_with_detected_encoding(path_obj):
    errors = []
    for encoding in _candidate_encodings():
        try:
            with path_obj.open("r", encoding=encoding, newline="") as stream:
                reader = csv.DictReader(stream)
                return list(reader), reader.fieldnames or [], encoding
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
    raise ValueError(f"Mine.csv 编码无法读取: {'; '.join(errors)}")


def _validate_required_columns(fieldnames):
    missing = [name for name in REQUIRED_COLUMNS if name not in (fieldnames or [])]
    if missing:
        raise ValueError(f"Mine.csv 缺失关键字段: {', '.join(missing)}")


def _status_summary(statuses):
    values = sorted({item for item in statuses if item})
    return " / ".join(values) if values else None


def _build_subject_record(subject_code, rows):
    first_row = rows[0]
    city = _normalize_text(first_row.get("地市"))
    county = _normalize_text(first_row.get("区县"))
    location_text = _normalize_text(first_row.get("矿山位置"))
    warnings = []
    for row in rows[1:]:
        if _normalize_text(row.get("地市")) not in (None, city):
            warnings.append(f"{subject_code}: 地市存在不一致，按首条主值保留")
            break
    for row in rows[1:]:
        if _normalize_text(row.get("区县")) not in (None, county):
            warnings.append(f"{subject_code}: 区县存在不一致，按首条主值保留")
            break
    for row in rows[1:]:
        if _normalize_text(row.get("矿山位置")) not in (None, location_text):
            warnings.append(f"{subject_code}: 矿山位置存在不一致，按首条主值保留")
            break

    plots = []
    statuses = []
    start_years = []
    end_years = []
    for row in rows:
        plot = {
            "plot_code": _normalize_text(row.get("图斑编号")) or "",
            "restored_plot_code": _normalize_text(row.get("修复图斑编号")),
            "plot_category": _normalize_text(row.get("图斑小类")),
            "repair_status": _normalize_text(row.get("修复状态")),
            "repair_mode": _normalize_text(row.get("修复模式")),
            "completed_at": _normalize_text(row.get("完成时间")),
            "area": _parse_float(row.get("Area"), field_name="Area"),
            "untreated_area": _parse_float(row.get("未治理面积"), field_name="未治理面积"),
            "center_lng": _parse_float(row.get("中心经度") or row.get("Lon"), field_name="经度"),
            "center_lat": _parse_float(row.get("中心纬度") or row.get("Lat"), field_name="纬度"),
            "closed_year": _normalize_text(row.get("关闭年度")),
            "plot_attr": _normalize_text(row.get("图斑属性")),
            "city": city,
            "county": county,
            "location_text": location_text,
        }
        plots.append(plot)
        if plot["repair_status"]:
            statuses.append(plot["repair_status"])
        closed_year = _extract_year(plot["closed_year"])
        completed_year = _extract_year(row.get("填报日期")) or _extract_year(plot["completed_at"])
        if closed_year:
            start_years.append(closed_year)
        if completed_year:
            end_years.append(completed_year)
        elif closed_year:
            end_years.append(closed_year)
    return {
        "tbbh": subject_code,
        "subject_code": subject_code,
        "city": city,
        "county": county,
        "location_text": location_text,
        "plot_count": len(plots),
        "area_total": sum(item["area"] for item in plots),
        "untreated_area_total": sum(item["untreated_area"] for item in plots),
        "status_summary": _status_summary(statuses),
        "plots": plots,
        "start_year": min(start_years) if start_years else None,
        "end_year": max(end_years) if end_years else None,
        "warnings": warnings,
    }


def _load_csv_rows(csv_path):
    path_obj = Path(csv_path).expanduser().resolve()
    if not path_obj.is_file():
        raise ValueError(f"Mine.csv 不存在: {path_obj}")
    rows, fieldnames, _encoding = _read_csv_with_detected_encoding(path_obj)
    _validate_required_columns(fieldnames)
    return path_obj, rows


def reset_project_domain_tables(commit=True):
    db.session.query(JiangxiMinePlot).delete()
    db.session.query(ProjectBackupRecord).delete()
    db.session.query(ProjectExportRecord).delete()
    db.session.query(ProjectActivityLog).delete()
    db.session.query(ProjectDataset).delete()
    db.session.query(ProjectMineBinding).delete()
    db.session.query(Project).delete()
    db.session.query(Analysis).delete()
    if commit:
        db.session.commit()
    else:
        db.session.flush()


def get_jiangxi_seed_database_state():
    default_project = Project.query.filter_by(
        remark=DEFAULT_JIANGXI_REMARK,
        deleted_at=None,
    ).first()
    if default_project:
        return "ready"

    domain_models = (
        Project,
        ProjectDataset,
        ProjectMineBinding,
        JiangxiMinePlot,
        ProjectActivityLog,
        ProjectBackupRecord,
        ProjectExportRecord,
        Analysis,
    )
    if any(model.query.count() for model in domain_models):
        return "invalid"
    return "empty"


def aggregate_subject_rows(csv_path):
    _path_obj, rows = _load_csv_rows(csv_path)
    grouped = defaultdict(list)
    skipped = 0
    warnings = []
    for row in rows:
        subject_code = _normalize_text(row.get("主体编号"))
        if not subject_code:
            skipped += 1
            warnings.append("存在主体编号为空的记录，已跳过")
            continue
        try:
            _parse_float(row.get("Area"), field_name="Area")
            _parse_float(row.get("未治理面积"), field_name="未治理面积")
            _parse_float(row.get("中心经度") or row.get("Lon"), field_name="经度")
            _parse_float(row.get("中心纬度") or row.get("Lat"), field_name="纬度")
        except ValueError as exc:
            skipped += 1
            warnings.append(f"{subject_code}: {exc}，已跳过该行")
            continue
        grouped[subject_code].append(row)
    sorted_items = sorted(
        grouped.items(),
        key=lambda item: (
            _normalize_text(item[1][0].get("地市")) or "",
            _normalize_text(item[1][0].get("区县")) or "",
            item[0],
        ),
    )
    subjects = []
    for subject_code, subject_rows in sorted_items:
        subject = _build_subject_record(subject_code, subject_rows)
        warnings.extend(subject["warnings"])
        subjects.append(subject)
    return subjects, skipped, warnings, len(rows)


def _read_workbook_rows(workbook_path):
    path_obj = Path(workbook_path).expanduser().resolve()
    if not path_obj.is_file():
        raise ValueError(f"江西图斑工作簿不存在: {path_obj}")
    workbook = load_workbook(path_obj, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = [str(value or "").strip() for value in next(worksheet.iter_rows(values_only=True))]
        missing = [name for name in WORKBOOK_REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ValueError(f"江西图斑工作簿缺少关键字段: {', '.join(missing)}")
        return [dict(zip(headers, row)) for row in worksheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()


def build_map_aligned_records(workbook_path, tbbh_to_map_fid=None):
    records = []
    seen_excel_fids = set()
    seen_tbbh = set()
    tbbh_to_map_fid = tbbh_to_map_fid or {}
    for row_number, row in enumerate(_read_workbook_rows(workbook_path), start=2):
        raw_excel_fid = row.get("FID")
        try:
            raw_excel_fid_text = "" if raw_excel_fid is None else str(raw_excel_fid).strip()
            if not raw_excel_fid_text:
                raise ValueError("empty FID")
            excel_fid_value = float(raw_excel_fid_text)
            if not math.isfinite(excel_fid_value) or not excel_fid_value.is_integer():
                raise ValueError("FID must be a finite integer")
            excel_fid = int(excel_fid_value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"第 {row_number} 行 Excel FID 非法: {raw_excel_fid}") from exc
        if excel_fid < 0:
            raise ValueError(f"第 {row_number} 行 Excel FID 必须为非负整数: {excel_fid}")
        tbbh = normalize_tbbh(row.get("TBBH"))
        if excel_fid in seen_excel_fids or tbbh in seen_tbbh:
            raise ValueError(f"第 {row_number} 行 Excel FID 或 TBBH 重复: {excel_fid} / {tbbh}")
        seen_excel_fids.add(excel_fid)
        seen_tbbh.add(tbbh)
        records.append(
            {
                "tbbh": tbbh,
                "excel_fid": excel_fid,
                "map_fid": tbbh_to_map_fid.get(tbbh),
                "mine_name": _normalize_text(row.get("矿山位置_")) or f"图斑 {tbbh}",
                "city": _normalize_text(row.get("SHI")) or _normalize_text(row.get("地市")),
                "county": _normalize_text(row.get("XIAN")) or _normalize_text(row.get("区县")),
                "area_hectare": _parse_float(row.get("面积"), field_name="面积") / 10000.0,
                "repair_status": _normalize_text(row.get("修复状态")),
                "repair_mode": _normalize_text(row.get("修复模式")),
                "completed_at": _normalize_text(row.get("完成时间")),
                "plot_category": _normalize_text(row.get("图斑小类")),
                "untreated_area": _parse_float(row.get("未治理面积"), field_name="未治理面积") / 10000.0,
                "center_lng": _parse_float(row.get("中心经度_"), field_name="中心经度"),
                "center_lat": _parse_float(row.get("中心纬度_"), field_name="中心纬度"),
            }
        )
    if not records:
        raise ValueError("江西图斑工作簿没有有效记录")
    if tbbh_to_map_fid and set(tbbh_to_map_fid) != seen_tbbh:
        raise ValueError("Excel TBBH 集合与权威资产 manifest 不一致")
    return sorted(records, key=lambda item: item["tbbh"])


def _replace_project_with_map_aligned_records(project, records):
    expected_tbbhs = {record["tbbh"] for record in records}
    orphaned_datasets = [
        dataset.tbbh
        for dataset in ProjectDataset.query.filter_by(project_id=project.id).all()
        if dataset.tbbh and dataset.tbbh not in expected_tbbhs
    ]
    if orphaned_datasets:
        raise ValueError(
            "同步会产生没有项目绑定的数据集: " + ", ".join(sorted(set(orphaned_datasets)))
        )
    db.session.query(JiangxiMinePlot).filter_by(project_id=project.id).delete(synchronize_session="fetch")
    db.session.query(ProjectMineBinding).filter_by(project_id=project.id).delete(synchronize_session="fetch")
    db.session.flush()
    for sort_order, record in enumerate(records, start=1):
        db.session.add(
            ProjectMineBinding(
                project_id=project.id,
                tbbh=record["tbbh"],
                mine_name_snapshot=record["mine_name"],
                city_snapshot=record["city"],
                area_snapshot=record["area_hectare"],
                status_snapshot=record["repair_status"],
                sort_order=sort_order,
            )
        )
        db.session.add(
            JiangxiMinePlot(
                project_id=project.id,
                tbbh=record["tbbh"],
                city=record["city"],
                county=record["county"],
                location_text=record["mine_name"],
                plot_code=record["tbbh"],
                restored_plot_code=record["tbbh"],
                plot_category=record["plot_category"],
                repair_status=record["repair_status"],
                repair_mode=record["repair_mode"],
                completed_at=record["completed_at"],
                area=record["area_hectare"],
                untreated_area=record["untreated_area"],
                center_lng=record["center_lng"],
                center_lat=record["center_lat"],
            )
        )
    project.monitor_start_year = WORKBOOK_MONITOR_START_YEAR
    project.monitor_end_year = WORKBOOK_MONITOR_END_YEAR


def _read_manifest_mapping(manifest_path):
    manifest_path = Path(manifest_path).expanduser().resolve()
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    if data.get("status") != "ok":
        raise ValueError("江西资产 manifest 状态不是 ok")
    mapping = ((data.get("mapping") or {}).get("tbbh_to_map_fid") or {})
    if not mapping:
        raise ValueError("江西资产 manifest 缺少 TBBH 映射")
    normalized = {}
    reverse = {}
    for key, value in mapping.items():
        tbbh = normalize_tbbh(key)
        try:
            map_fid = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"江西资产 manifest map_fid 非法: {value}") from exc
        if map_fid <= 0 or tbbh in normalized or map_fid in reverse:
            raise ValueError(f"江西资产 manifest TBBH/map_fid 重复或非法: {tbbh}/{map_fid}")
        normalized[tbbh] = map_fid
        reverse[map_fid] = tbbh
    return normalized


def _manifest_sha256(manifest_path):
    digest = hashlib.sha256()
    with Path(manifest_path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sync_jiangxi_project_from_workbook(workbook_path, actor="system", manifest_path=None):
    if not manifest_path:
        raise ValueError("江西同步必须提供资产 manifest")
    manifest_sha256 = _manifest_sha256(manifest_path)
    tbbh_to_map_fid = _read_manifest_mapping(manifest_path)
    records = build_map_aligned_records(workbook_path, tbbh_to_map_fid)
    project = Project.query.filter_by(remark=DEFAULT_JIANGXI_REMARK, deleted_at=None).first()
    if not project:
        raise ValueError("未找到可对齐的江西默认项目")
    expected_tbbhs = {record["tbbh"] for record in records}
    existing_tbbhs = {row.tbbh for row in ProjectMineBinding.query.filter_by(project_id=project.id).all()}
    existing_plots = JiangxiMinePlot.query.filter_by(project_id=project.id).all()
    expected_rows = {_plot_content_signature({
        **record,
        "location_text": record["mine_name"],
        "plot_code": record["tbbh"],
        "restored_plot_code": record["tbbh"],
        "area": record["area_hectare"],
    }) for record in records}
    existing_rows = {_plot_content_signature(row) for row in existing_plots}
    expected_bindings = {
        (
            record["tbbh"],
            record["mine_name"],
            record["city"],
            record["area_hectare"],
            record["repair_status"],
            sort_order,
        )
        for sort_order, record in enumerate(records, start=1)
    }
    existing_bindings = {
        (
            row.tbbh,
            row.mine_name_snapshot,
            row.city_snapshot,
            row.area_snapshot,
            row.status_snapshot,
            row.sort_order,
        )
        for row in ProjectMineBinding.query.filter_by(project_id=project.id).all()
    }
    if (
        existing_tbbhs == expected_tbbhs
        and existing_rows == expected_rows
        and existing_bindings == expected_bindings
        and project.monitor_start_year == WORKBOOK_MONITOR_START_YEAR
        and project.monitor_end_year == WORKBOOK_MONITOR_END_YEAR
    ):
        return {
            "project_id": project.id,
            "tbbh_count": len(records),
            "mine_count": len(records),
            "plot_count": len(records),
            "asset_manifest_sha256": manifest_sha256,
            "changed": False,
        }
    try:
        _replace_project_with_map_aligned_records(project, records)
        db.session.add(
            ProjectActivityLog(
                project_id=project.id,
                event_type="jiangxi_tbbh_synced",
                actor=actor,
                payload_json=json.dumps({"source": Path(workbook_path).name, "tbbh_count": len(records)}, ensure_ascii=False),
            )
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {
        "project_id": project.id,
        "tbbh_count": len(records),
        "mine_count": len(records),
        "plot_count": len(records),
        "asset_manifest_sha256": manifest_sha256,
        "changed": True,
    }


def seed_jiangxi_project_from_csv(csv_path, actor="system", workbook_path=None, manifest_path=None):
    if not workbook_path:
        raise ValueError("江西发布种子必须提供权威 Excel；Mine.csv 仅用于校验，不能单独建库")
    if not manifest_path:
        raise ValueError("江西发布种子必须提供资产 manifest")
    csv_path_obj, _rows = _load_csv_rows(csv_path)
    tbbh_to_map_fid = _read_manifest_mapping(manifest_path)
    records = build_map_aligned_records(workbook_path, tbbh_to_map_fid)
    try:
        reset_project_domain_tables(commit=False)
        project = Project(
            name=DEFAULT_JIANGXI_PROJECT_NAME,
            region=DEFAULT_JIANGXI_REGION,
            manager="admin",
            remark=DEFAULT_JIANGXI_REMARK,
            status="active",
            monitor_start_year=WORKBOOK_MONITOR_START_YEAR,
            monitor_end_year=WORKBOOK_MONITOR_END_YEAR,
        )
        db.session.add(project)
        db.session.flush()
        _replace_project_with_map_aligned_records(project, records)
        payload = {
            "source": Path(workbook_path).name,
            "csv_validation_source": str(csv_path_obj),
            "tbbh_count": len(records),
            "plot_count": len(records),
        }
        payload["asset_manifest_sha256"] = _manifest_sha256(manifest_path)
        db.session.add(
            ProjectActivityLog(
                project_id=project.id,
                event_type="jiangxi_seed_completed",
                actor=actor,
                payload_json=json.dumps(payload, ensure_ascii=False),
            )
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {
        "project_id": project.id,
        "project_name": project.name,
        "project_count": 1,
        "tbbh_count": len(records),
        "subject_count": len(records),
        "plot_count": len(records),
        "skipped_count": 0,
        "monitor_start_year": WORKBOOK_MONITOR_START_YEAR,
        "monitor_end_year": WORKBOOK_MONITOR_END_YEAR,
        "asset_manifest_sha256": _manifest_sha256(manifest_path),
    }

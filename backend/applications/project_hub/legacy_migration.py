"""Import historical Jiangxi outputs without treating any legacy FID as a business key."""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import desc

from applications.extensions import db
from applications.kml_roi.index_sync import INDEX_FILE_MAP, sync_miner_index_rows
from applications.models.analysis import Analysis
from applications.models.project import Project, ProjectActivityLog, ProjectDataset, ProjectMineBinding
from applications.project_hub.tbbh_identity import normalize_tbbh
from applications.region_source import load_kml_root, resolve_default_jiangxi_kmz

KML_NS = {"kml": "http://www.opengis.net/kml/2.2"}
YEAR_RE = re.compile(r"(20\d{2})")
MASK_YEAR_RE = re.compile(r"\+(20\d{2})_mask\.png$")
LEGACY_PROJECT_REMARK = "legacy_migration:auto"
MAX_DB_INT = 2147483647


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _default_miner_root() -> Path:
    return _repo_root() / "miner"


def _default_output_root() -> Path:
    return Path(
        os.getenv("MINER_CHANGE_OUTPUT_ROOT") or _default_miner_root() / "change_matrix_outputs"
    )


def _default_kml_path() -> Path:
    return resolve_default_jiangxi_kmz()


def _default_static_root() -> Path:
    return _repo_root() / "backend" / "static"


def _default_manifest_path() -> Path:
    return _repo_root() / "docker" / "standalone" / "runtime_data" / "Jiangxi_asset_manifest.json"


def _safe_json_load(value, default=None):
    if value in (None, ""):
        return {} if default is None else default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        return {} if default is None else default


def _json_dump(value):
    return json.dumps(value or {}, ensure_ascii=False)


def _to_int(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        number_float = float(text)
    except Exception:
        return None
    if not math.isfinite(number_float) or not number_float.is_integer():
        return None
    number = int(number_float)
    return number if 0 < number <= MAX_DB_INT else None


def _to_float(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return float(text)
    except Exception:
        return None


def _sha256_file(path_obj: Path):
    digest = hashlib.sha256()
    with path_obj.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_manifest_mapping(manifest_path: Path):
    if not manifest_path.is_file():
        raise ValueError(f"历史迁移必须提供江西资产 manifest: {manifest_path}")
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ValueError(f"江西资产 manifest 无法读取: {manifest_path}: {exc}") from exc
    if data.get("status") != "ok":
        raise ValueError("江西资产 manifest 状态不是 ok，禁止迁移历史数据")
    raw_mapping = (data.get("mapping") or {}).get("tbbh_to_map_fid") or {}
    if not raw_mapping:
        raise ValueError("江西资产 manifest 缺少 TBBH/map_fid 映射")
    tbbh_to_map_fid = {}
    map_fid_to_tbbh = {}
    for raw_tbbh, raw_map_fid in raw_mapping.items():
        tbbh = normalize_tbbh(raw_tbbh)
        map_fid = _to_int(raw_map_fid)
        if map_fid is None or tbbh in tbbh_to_map_fid or map_fid in map_fid_to_tbbh:
            raise ValueError(f"manifest 中 TBBH/map_fid 重复或非法: {raw_tbbh}/{raw_map_fid}")
        tbbh_to_map_fid[tbbh] = map_fid
        map_fid_to_tbbh[map_fid] = tbbh
    return tbbh_to_map_fid, map_fid_to_tbbh


def _resolve_legacy_map_fid(raw_value, map_fid_to_tbbh, *, source):
    map_fid = _to_int(raw_value)
    if map_fid is None:
        raise ValueError(f"{source} 的历史编号不是正整数: {raw_value}")
    if map_fid not in map_fid_to_tbbh:
        raise ValueError(
            f"{source} 的历史编号无法唯一映射为 TBBH: map_fid={map_fid}；请声明正确的 FID 命名空间"
        )
    return map_fid, map_fid_to_tbbh[map_fid]


def _find_or_create_project(project_name: str, manager: str, min_year, max_year):
    project = Project.query.filter_by(name=project_name, deleted_at=None).first()
    project_created = False
    if project is None:
        project = Project(
            name=project_name,
            region="江西省",
            manager=manager or "admin",
            remark=LEGACY_PROJECT_REMARK,
            status="active",
            monitor_start_year=min_year,
            monitor_end_year=max_year,
        )
        db.session.add(project)
        db.session.flush()
        project_created = True
        return project, project_created

    project.deleted_at = None
    project.region = project.region or "江西省"
    project.manager = project.manager or manager or "admin"
    project.status = "active" if project.status == "draft" else project.status
    project.remark = project.remark or LEGACY_PROJECT_REMARK
    if min_year is not None and (project.monitor_start_year is None or min_year < project.monitor_start_year):
        project.monitor_start_year = min_year
    if max_year is not None and (project.monitor_end_year is None or max_year > project.monitor_end_year):
        project.monitor_end_year = max_year
    return project, project_created


def _load_kml_snapshots(kml_path: Path, map_fid_to_tbbh):
    snapshots = {}
    if not kml_path.exists():
        return snapshots
    root = load_kml_root(kml_path)
    for placemark in root.findall(".//kml:Placemark", KML_NS):
        fields = {
            str(item.attrib.get("name") or "").strip(): "".join(item.itertext()).strip()
            for item in placemark.findall(".//kml:SimpleData", KML_NS)
        }
        raw_tbbh = fields.get("TBBH") or fields.get("tbbh")
        if raw_tbbh:
            tbbh = normalize_tbbh(raw_tbbh)
            map_fid = _to_int(fields.get("FID_1") or fields.get("序号"))
            if map_fid is not None and map_fid_to_tbbh.get(map_fid) != tbbh:
                raise ValueError(f"KML TBBH/map_fid 冲突: {tbbh}/{map_fid}")
        else:
            map_fid, tbbh = _resolve_legacy_map_fid(
                fields.get("FID_1")
                or fields.get("序号")
                or placemark.findtext("kml:name", default="", namespaces=KML_NS),
                map_fid_to_tbbh,
                source=f"KML {placemark.findtext('kml:name', default='', namespaces=KML_NS)}",
            )
        mine_name = (
            fields.get("GGKSMC")
            or fields.get("SBKSMC")
            or fields.get("ZLKSMC")
            or placemark.findtext("kml:name", default="", namespaces=KML_NS).strip()
            or f"矿山 {map_fid}"
        )
        if tbbh in snapshots:
            raise ValueError(f"KML 存在重复 TBBH: {tbbh}")
        snapshots[tbbh] = {
            "mine_name_snapshot": mine_name,
            "city_snapshot": fields.get("SHI") or fields.get("SHI_1") or "",
            "area_snapshot": _to_float(fields.get("TBTYMJ_1") or fields.get("TBTYMJ") or fields.get("SHAPE_Area")),
            "status_snapshot": fields.get("HFZLQK") or fields.get("ZLHFZLQK") or "",
            "map_fid": map_fid,
        }
    return snapshots


def _iter_change_output_entries(output_root: Path, map_fid_to_tbbh):
    entries = []
    if not output_root.exists():
        return entries
    for child in sorted(output_root.iterdir(), key=lambda item: item.name):
        if not child.is_dir():
            continue
        map_fid, tbbh = _resolve_legacy_map_fid(child.name, map_fid_to_tbbh, source=f"历史结果目录 {child}")
        file_names = sorted(item.name for item in child.iterdir() if item.is_file())
        if not file_names:
            continue
        years = sorted(
            {
                int(match.group(1))
                for name in file_names
                for match in [MASK_YEAR_RE.search(name)]
                if match
            }
        )
        entries.append(
            {
                "tbbh": tbbh,
                "map_fid": map_fid,
                "path": child,
                "available_years": years,
                "file_names": file_names,
                "has_change_matrix_percent": "change_matrix_percent_rownorm.csv" in file_names,
                "has_class_ratio": "class_ratio_percent.json" in file_names,
            }
        )
    return entries


def _workbook_summary(path_obj: Path):
    if not path_obj.exists():
        return {"path": path_obj, "map_fids": [], "available_years": []}
    workbook = load_workbook(path_obj, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        years = sorted(
            {int(match.group(1)) for value in header_row for match in [YEAR_RE.search(str(value or ""))] if match}
        )
        map_fids = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            raw_map_fid = row[0] if row else None
            if raw_map_fid in (None, ""):
                continue
            map_fid = _to_int(raw_map_fid)
            if map_fid is None:
                raise ValueError(f"指数工作簿 {path_obj} 第 {row_number} 行 map_fid 非法: {raw_map_fid}")
            map_fids.append(map_fid)
        if len(map_fids) != len(set(map_fids)):
            raise ValueError(f"指数工作簿 {path_obj} 存在重复 map_fid")
        return {"path": path_obj, "map_fids": sorted(set(map_fids)), "available_years": years}
    finally:
        workbook.close()


def _resolve_analysis_file_path(raw_path: str, static_root: Path):
    text = str(raw_path or "").strip()
    if not text:
        return text
    if text.startswith("/_uploads/photos/"):
        candidate = static_root / "upload" / text[len("/_uploads/photos/"):]
        return str(candidate) if candidate.exists() else text
    if text.startswith("static/"):
        candidate = static_root.parent / text
        return str(candidate) if candidate.exists() else text
    return text


def _load_analysis_rows():
    return Analysis.query.order_by(desc(Analysis.create_time), desc(Analysis.id)).all()


def _sync_analysis_workbooks(analysis_rows, miner_root: Path):
    generated_files = []
    synced_rows = 0
    skipped_existing = 0
    warnings = []
    for row in analysis_rows:
        if row.type != 8:
            continue
        meta = _safe_json_load(row.data, {})
        index_type = str(meta.get("index_type") or "").strip().upper()
        year = str(meta.get("year") or "").strip()
        fid_stats = meta.get("map_fid_stats") or meta.get("fid_stats") or []
        if not index_type or not fid_stats:
            continue
        target_path = miner_root / INDEX_FILE_MAP.get(index_type, "")
        existed_before = target_path.exists()
        normalized_rows = [
            {
                "fid": item.get("map_fid", item.get("fid")),
                "mean": item.get("mean"),
            }
            for item in fid_stats
            if item.get("map_fid", item.get("fid")) not in (None, "")
            and item.get("mean") is not None
        ]
        result = sync_miner_index_rows(index_type, year, normalized_rows, miner_dir=miner_root, overwrite_existing=False)
        if result.get("synced"):
            synced_rows += int(result.get("rows") or 0)
            skipped_existing += int(result.get("skipped_existing") or 0)
            if not existed_before and index_type not in generated_files:
                generated_files.append(index_type)
        else:
            warnings.append({"analysis_id": row.id, "index_type": index_type, "reason": result.get("reason")})
    return {
        "generated_files": generated_files,
        "synced_rows": synced_rows,
        "skipped_existing": skipped_existing,
        "warnings": warnings,
    }


def _build_workbook_datasets(index_summaries):
    datasets = []
    for index_type, summary in sorted(index_summaries.items()):
        path_obj = summary["path"]
        if not path_obj.exists():
            continue
        years = summary["available_years"]
        datasets.append(
            {
                "dataset_kind": "report",
                "display_name": f"历史 {index_type} 指数时序",
                "file_path": str(path_obj),
                "source_format": "xlsx",
                "tbbh": None,
                "year_start": min(years) if years else None,
                "year_end": max(years) if years else None,
                "slice_config_json": {
                    "legacy_source": "miner_index_workbook",
                    "legacy_fid_namespace": "map_fid",
                    "index_type": index_type,
                    "map_fid_count": len(summary["map_fids"]),
                    "available_years": years,
                },
            }
        )
    return datasets


def _build_change_output_datasets(entries):
    datasets = []
    for entry in entries:
        datasets.append(
            {
                "dataset_kind": "inference_result",
                "display_name": f"历史变化矩阵输出 TBBH {entry['tbbh']}",
                "file_path": str(entry["path"]),
                "source_format": "change_matrix_dir",
                "tbbh": entry["tbbh"],
                "year_start": min(entry["available_years"]) if entry["available_years"] else None,
                "year_end": max(entry["available_years"]) if entry["available_years"] else None,
                "slice_config_json": {
                    "legacy_source": "change_matrix_outputs",
                    "legacy_fid_namespace": "map_fid",
                    "map_fid": entry["map_fid"],
                    "tbbh": entry["tbbh"],
                    "available_years": entry["available_years"],
                    "file_count": len(entry["file_names"]),
                    "has_change_matrix_percent": entry["has_change_matrix_percent"],
                    "has_class_ratio": entry["has_class_ratio"],
                },
            }
        )
    return datasets


def _analysis_tbbhs(meta, map_fid_to_tbbh, analysis_id):
    declared_tbbhs = meta.get("matched_tbbh_list") or []
    if declared_tbbhs:
        result = []
        for raw_tbbh in declared_tbbhs:
            tbbh = normalize_tbbh(raw_tbbh)
            if tbbh not in map_fid_to_tbbh.values():
                raise ValueError(f"分析记录 {analysis_id} 的 TBBH 不在权威 manifest: {tbbh}")
            if tbbh not in result:
                result.append(tbbh)
        return result
    raw_fids = meta.get("matched_fid_list") or []
    result = []
    for raw_fid in raw_fids:
        _map_fid, tbbh = _resolve_legacy_map_fid(raw_fid, map_fid_to_tbbh, source=f"分析记录 {analysis_id}")
        if tbbh not in result:
            result.append(tbbh)
    return result


def _build_analysis_datasets(analysis_rows, static_root: Path, map_fid_to_tbbh):
    datasets = []
    for row in analysis_rows:
        meta = _safe_json_load(row.data, {})
        matched_tbbhs = _analysis_tbbhs(meta, map_fid_to_tbbh, row.id)
        if row.type == 8:
            index_type = str(meta.get("index_type") or "INDEX").strip().upper()
            year = str(meta.get("year") or "").strip()
            display_name = f"历史 {index_type} 分析记录 #{row.id}"
            if year:
                display_name = f"历史 {index_type} 分析记录 {year} #{row.id}"
            year_start = _to_int(year)
            year_end = _to_int(year)
        else:
            index_type = ""
            display_name = f"历史分析记录 #{row.id}"
            year_start = year_end = None
        datasets.append(
            {
                "dataset_kind": "inference_result",
                "display_name": display_name,
                "file_path": _resolve_analysis_file_path(row.after_img or row.before_img, static_root),
                "source_format": "analysis_record",
                "tbbh": matched_tbbhs[0] if len(matched_tbbhs) == 1 else None,
                "year_start": year_start,
                "year_end": year_end,
                "slice_config_json": {
                    "legacy_source": "analysis_history",
                    "legacy_fid_namespace": "map_fid",
                    "analysis_id": row.id,
                    "analysis_type": row.type,
                    "index_type": index_type,
                    "matched_tbbh_list": matched_tbbhs[:200],
                    "raw_matched_map_fid_list": meta.get("matched_fid_list") or [],
                    "map_fid_stats": meta.get("map_fid_stats") or [],
                    "raw_after_img": row.after_img,
                    "raw_before_img": row.before_img,
                },
            }
        )
    return datasets


def _collect_tbbhs(change_entries, index_summaries, analysis_rows, map_fid_to_tbbh):
    tbbhs = {entry["tbbh"] for entry in change_entries}
    for summary in index_summaries.values():
        for map_fid in summary["map_fids"]:
            _map_fid, tbbh = _resolve_legacy_map_fid(map_fid, map_fid_to_tbbh, source=f"指数工作簿 {summary['path']}")
            tbbhs.add(tbbh)
    for row in analysis_rows:
        tbbhs.update(_analysis_tbbhs(_safe_json_load(row.data, {}), map_fid_to_tbbh, row.id))
    return sorted(tbbhs)


def _dataset_key(item):
    return (
        str(item.get("dataset_kind") or "").strip(),
        str(item.get("file_path") or "").strip(),
        str(item.get("tbbh") or "").strip(),
    )


def _upsert_mine_bindings(project: Project, tbbhs, snapshots):
    existing = {row.tbbh: row for row in project.mines}
    added = 0
    for sort_order, tbbh in enumerate(tbbhs, start=1):
        row = existing.get(tbbh)
        if row is None:
            row = ProjectMineBinding(project_id=project.id, tbbh=tbbh)
            db.session.add(row)
            existing[tbbh] = row
            added += 1
        snapshot = snapshots.get(tbbh, {})
        row.mine_name_snapshot = snapshot.get("mine_name_snapshot") or row.mine_name_snapshot or f"矿山 {tbbh}"
        row.city_snapshot = snapshot.get("city_snapshot") or row.city_snapshot
        if snapshot.get("area_snapshot") is not None:
            row.area_snapshot = snapshot["area_snapshot"]
        row.status_snapshot = snapshot.get("status_snapshot") or row.status_snapshot
        row.sort_order = sort_order
    return added, len(existing)


def _upsert_datasets(project: Project, dataset_items):
    existing = {_dataset_key(item.__dict__): item for item in project.datasets}
    added = 0
    for item in dataset_items:
        key = _dataset_key(item)
        row = existing.get(key)
        if row is None:
            row = ProjectDataset(
                project_id=project.id,
                dataset_kind=item["dataset_kind"],
                display_name=item["display_name"],
                file_path=item["file_path"],
                source_format=item["source_format"],
                tbbh=item.get("tbbh"),
                year_start=item.get("year_start"),
                year_end=item.get("year_end"),
                slice_config_json=_json_dump(item.get("slice_config_json")),
            )
            db.session.add(row)
            existing[key] = row
            added += 1
            continue
        row.display_name = item["display_name"]
        row.source_format = item["source_format"]
        row.tbbh = item.get("tbbh")
        row.year_start = item.get("year_start")
        row.year_end = item.get("year_end")
        row.slice_config_json = _json_dump(item.get("slice_config_json"))
    return added, len(existing)


def _append_activity(project_id, event_type, payload, actor):
    db.session.add(
        ProjectActivityLog(
            project_id=project_id,
            event_type=event_type,
            actor=actor or "system",
            payload_json=_json_dump(payload),
        )
    )


def _write_migration_report(path_obj, report):
    if not path_obj:
        return
    path_obj.parent.mkdir(parents=True, exist_ok=True)
    path_obj.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def migrate_legacy_project_data(
    project_name: str = "江西历史成果迁移项目",
    manager: str = "admin",
    miner_root=None,
    output_root=None,
    kml_path=None,
    static_root=None,
    asset_manifest_path=None,
    migration_report_path=None,
):
    miner_root = Path(miner_root) if miner_root else _default_miner_root()
    output_root = Path(output_root) if output_root else _default_output_root()
    kml_path = Path(kml_path) if kml_path else _default_kml_path()
    static_root = Path(static_root) if static_root else _default_static_root()
    manifest_path = Path(asset_manifest_path) if asset_manifest_path else _default_manifest_path()
    report_path = Path(migration_report_path) if migration_report_path else output_root / "legacy_migration_report.json"
    report = {"status": "failed", "manifest": str(manifest_path), "records": [], "unresolved": []}

    try:
        _tbbh_to_map_fid, map_fid_to_tbbh = _load_manifest_mapping(manifest_path)
        analysis_rows = _load_analysis_rows()
        index_summaries = {
            index_type: _workbook_summary(miner_root / filename)
            for index_type, filename in INDEX_FILE_MAP.items()
        }
        change_entries = _iter_change_output_entries(output_root, map_fid_to_tbbh)
        snapshots = _load_kml_snapshots(kml_path, map_fid_to_tbbh)
        tbbhs = _collect_tbbhs(change_entries, index_summaries, analysis_rows, map_fid_to_tbbh)
        index_sync = _sync_analysis_workbooks(analysis_rows, miner_root)

        all_years = [year for entry in change_entries for year in entry["available_years"]]
        all_years.extend(year for summary in index_summaries.values() for year in summary["available_years"])
        all_years.extend(
            year
            for row in analysis_rows
            for year in [_to_int(_safe_json_load(row.data, {}).get("year"))]
            if year is not None
        )
        min_year = min(all_years) if all_years else None
        max_year = max(all_years) if all_years else None

        project, project_created = _find_or_create_project(project_name, manager, min_year, max_year)
        mine_bindings_added, mine_count = _upsert_mine_bindings(project, tbbhs, snapshots)
        dataset_items = _build_workbook_datasets(index_summaries)
        dataset_items.extend(_build_change_output_datasets(change_entries))
        dataset_items.extend(_build_analysis_datasets(analysis_rows, static_root, map_fid_to_tbbh))
        datasets_added, dataset_count = _upsert_datasets(project, dataset_items)

        for item in dataset_items:
            item_path = Path(item["file_path"])
            report["records"].append(
                {
                    "source_path": str(item_path),
                    "tbbh": item.get("tbbh"),
                    "map_fid": (
                        (item.get("slice_config_json") or {}).get("map_fid")
                        or (snapshots.get(item.get("tbbh"), {}) or {}).get("map_fid")
                        or (
                            _tbbh_to_map_fid.get(item.get("tbbh"))
                            if item.get("tbbh")
                            else None
                        )
                    ),
                    "sha256": _sha256_file(item_path) if item_path.is_file() else None,
                    "imported_at": datetime.now(timezone.utc).isoformat(),
                    "status": "imported",
                }
            )

        payload = {
            "project_created": project_created,
            "mine_bindings_added": mine_bindings_added,
            "mine_count": mine_count,
            "datasets_added": datasets_added,
            "dataset_count": dataset_count,
            "index_sync": index_sync,
            "change_output_count": len(change_entries),
            "analysis_record_count": len(analysis_rows),
            "asset_manifest_sha256": _sha256_file(manifest_path),
            "migration_report_path": str(report_path),
        }
        _append_activity(project.id, "legacy_data_migrated", payload, actor=manager or "system")
        db.session.commit()

        report.update({"status": "ok", "project_id": project.id, "counts": payload})
        _write_migration_report(report_path, report)
        return {
            "project_id": project.id,
            "project_created": project_created,
            "mine_bindings_added": mine_bindings_added,
            "mine_count": mine_count,
            "datasets_added": datasets_added,
            "dataset_count": dataset_count,
            "index_sync": index_sync,
            "year_range": [min_year, max_year],
            "asset_manifest_sha256": payload["asset_manifest_sha256"],
            "migration_report_path": str(report_path),
        }
    except Exception as exc:
        db.session.rollback()
        report["error"] = str(exc)
        _write_migration_report(report_path, report)
        raise

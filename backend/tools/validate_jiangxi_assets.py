from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import struct
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

try:
    from applications.project_hub.tbbh_identity import (
        build_tbbh_map_fid_mapping,
        _parse_map_fid,
        normalize_tbbh,
        validate_tbbh_sets,
    )
except ModuleNotFoundError as exc:
    if exc.name != "flask":
        raise
    import importlib.util

    _identity_path = Path(__file__).resolve().parents[1] / "applications" / "project_hub" / "tbbh_identity.py"
    _identity_spec = importlib.util.spec_from_file_location("jiangxi_tbbh_identity_standalone", _identity_path)
    _identity_module = importlib.util.module_from_spec(_identity_spec)
    _identity_spec.loader.exec_module(_identity_module)
    build_tbbh_map_fid_mapping = _identity_module.build_tbbh_map_fid_mapping
    _parse_map_fid = _identity_module._parse_map_fid
    normalize_tbbh = _identity_module.normalize_tbbh
    validate_tbbh_sets = _identity_module.validate_tbbh_sets


KML_NS = {"kml": "http://www.opengis.net/kml/2.2"}
REQUIRED_CSV_COLUMNS = ("主体编号", "地市", "区县", "矿山位置", "Area")
MAP_FID_FIELDS = ("map_fid", "FID_1", "FID", "序号", "No", "OBJECTID")
# 只有明确的权威 TBBH 字段可以参与跨源身份关联；Mine.csv 的“主体编号”和
# 图斑业务字段不能被隐式当作 TBBH。
TBBH_FIELDS = ("TBBH", "tbbh")


def _repo_root():
    return Path(__file__).resolve().parents[2]


def default_asset_paths():
    runtime = _repo_root() / "docker" / "standalone" / "runtime_data"
    return {
        "excel": _repo_root() / "miner" / "data" / "348图斑_TableMERNet无图像预测结果.xlsx",
        "shp": runtime / "348个图斑.shp",
        "kmz": runtime / "Jiangxi_NaturalMine.kmz",
        "csv": runtime / "Mine.csv",
    }


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _file_info(path, count=None):
    return {
        "path": str(path),
        "sha256": sha256_file(path),
        "size": path.stat().st_size,
        "count": count,
    }


def _header_alias(fields, aliases):
    for alias in aliases:
        if alias in fields:
            return alias
    return None


def _to_number(value):
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _read_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(values, ())]
        rows = [dict(zip(headers, row)) for row in values]
    finally:
        workbook.close()

    tbbh_field = _header_alias(headers, TBBH_FIELDS[:2])
    if not tbbh_field:
        raise ValueError("Excel 缺少 TBBH 字段")
    map_fid_field = _header_alias(headers, ("map_fid", "序号", "FID_1"))
    normalized_rows = []
    for row_number, row in enumerate(rows, start=2):
        try:
            tbbh = normalize_tbbh(row.get(tbbh_field))
        except ValueError as exc:
            raise ValueError(f"Excel 第 {row_number} 行 TBBH 无效: {exc}") from exc
        normalized_rows.append({
            "tbbh": tbbh,
            "map_fid": row.get(map_fid_field) if map_fid_field else None,
            "excel_fid": row.get("FID"),
            "row": row_number,
        })
    mapping = (
        build_tbbh_map_fid_mapping(normalized_rows, "tbbh", "map_fid")
        if map_fid_field
        else {}
    )
    return {
        "rows": normalized_rows,
        "tbbh": [row["tbbh"] for row in normalized_rows],
        "mapping": mapping,
        "count": len(normalized_rows),
    }


def _read_dbf(path):
    raw = path.read_bytes()
    if len(raw) < 32:
        raise ValueError(f"DBF 文件过小: {path}")
    record_count = struct.unpack_from("<I", raw, 4)[0]
    header_length = struct.unpack_from("<H", raw, 8)[0]
    record_length = struct.unpack_from("<H", raw, 10)[0]
    fields = []
    offset = 32
    while offset + 32 <= len(raw) and raw[offset] != 0x0D:
        encoded_name = raw[offset : offset + 11].split(b"\0", 1)[0]
        try:
            name = encoded_name.decode("gb18030").strip()
        except UnicodeDecodeError:
            name = encoded_name.decode("latin1").strip()
        field_type = chr(raw[offset + 11])
        field_length = raw[offset + 16]
        fields.append((name, field_type, field_length))
        offset += 32
    if header_length > len(raw) or record_length <= 0:
        raise ValueError(f"DBF 头部无效: {path}")
    rows = []
    data_offset = header_length
    for index in range(record_count):
        start = data_offset + index * record_length
        end = start + record_length
        if end > len(raw):
            raise ValueError(f"DBF 记录不完整: {path} 第 {index + 1} 条")
        record = raw[start:end]
        values = {}
        cursor = 1
        for name, field_type, field_length in fields:
            value = record[cursor : cursor + field_length]
            cursor += field_length
            try:
                text = value.decode("gb18030").strip()
            except UnicodeDecodeError:
                text = value.decode("latin1").strip()
            if field_type in "NFI" and text:
                try:
                    text = float(text) if "." in text else int(text)
                except ValueError:
                    pass
            values[name] = text
        rows.append(values)
    return fields, rows


def _count_shp_records(path):
    raw = path.read_bytes()
    if len(raw) < 100:
        raise ValueError(f"SHP 文件过小: {path}")
    declared_words = struct.unpack_from(">i", raw, 24)[0]
    if declared_words * 2 > len(raw):
        raise ValueError(f"SHP 文件长度声明超过实际长度: {path}")
    count = 0
    offset = 100
    while offset < len(raw):
        if offset + 8 > len(raw):
            raise ValueError(f"SHP 记录头不完整: {path}")
        content_words = struct.unpack_from(">i", raw, offset + 4)[0]
        record_bytes = 8 + content_words * 2
        if content_words < 0 or offset + record_bytes > len(raw):
            raise ValueError(f"SHP 记录长度无效: {path} 第 {count + 1} 条")
        count += 1
        offset += record_bytes
    return count


def _read_shp(path, excel_mapping):
    sidecars = {
        suffix: path.with_suffix(suffix)
        for suffix in (".shx", ".dbf", ".prj")
    }
    missing = [str(item) for item in sidecars.values() if not item.is_file()]
    if missing:
        raise ValueError("SHP 缺少配套文件: " + ", ".join(missing))
    shp_count = _count_shp_records(path)
    fields, rows = _read_dbf(sidecars[".dbf"])
    if shp_count != len(rows):
        raise ValueError(f"SHP/DBF 记录数不一致: SHP={shp_count}, DBF={len(rows)}")
    tbbh_field = _header_alias([item[0] for item in fields], TBBH_FIELDS)
    map_fid_field = _header_alias([item[0] for item in fields], MAP_FID_FIELDS)
    records = []
    for row in rows:
        raw_map_fid = row.get(map_fid_field) if map_fid_field else None
        map_fid_text = str(raw_map_fid).strip() if raw_map_fid is not None else ""
        map_fid = _parse_map_fid(raw_map_fid) if map_fid_text else None
        raw_tbbh = row.get(tbbh_field) if tbbh_field else None
        tbbh = normalize_tbbh(raw_tbbh) if raw_tbbh not in (None, "") else None
        if map_fid is None and tbbh in excel_mapping:
            map_fid = excel_mapping[tbbh]
        if tbbh is None and map_fid in set(excel_mapping.values()):
            tbbh = next(key for key, value in excel_mapping.items() if value == map_fid)
        records.append({"tbbh": tbbh, "map_fid": map_fid})
    return {
        "records": records,
        "tbbh": [item["tbbh"] for item in records if item["tbbh"]],
        "count": shp_count,
        "files": sidecars,
    }


def _read_kmz(path, excel_mapping):
    with zipfile.ZipFile(path) as archive:
        kml_names = [name for name in archive.namelist() if name.lower().endswith(".kml")]
        if not kml_names:
            raise ValueError(f"KMZ 中没有 KML: {path}")
        root = ET.fromstring(archive.read(kml_names[0]))
    records = []
    for placemark in root.findall(".//kml:Placemark", KML_NS):
        fields = {}
        for item in placemark.findall(".//kml:SimpleData", KML_NS):
            fields[str(item.attrib.get("name") or "").strip()] = "".join(item.itertext()).strip()
        map_fid = None
        for field in MAP_FID_FIELDS:
            if str(fields.get(field, "")).strip():
                try:
                    map_fid = _parse_map_fid(fields[field])
                    break
                except ValueError:
                    raise
        if map_fid is None:
            name = placemark.findtext("kml:name", default="", namespaces=KML_NS).strip()
            if re.fullmatch(r"\d+", name):
                map_fid = _parse_map_fid(name)
        raw_tbbh = next((fields.get(field) for field in TBBH_FIELDS if fields.get(field)), None)
        tbbh = normalize_tbbh(raw_tbbh) if raw_tbbh else None
        if map_fid is None and tbbh in excel_mapping:
            map_fid = excel_mapping[tbbh]
        if tbbh is None and map_fid in set(excel_mapping.values()):
            tbbh = next(key for key, value in excel_mapping.items() if value == map_fid)
        records.append({"tbbh": tbbh, "map_fid": map_fid})
    return {
        "records": records,
        "tbbh": [item["tbbh"] for item in records if item["tbbh"]],
        "count": len(records),
    }


def _read_csv(path):
    errors = []
    rows = None
    encoding = None
    for candidate in ("gb18030", "gbk", "utf-8-sig", "utf-8"):
        try:
            with path.open("r", encoding=candidate, newline="") as stream:
                reader = csv.DictReader(stream)
                rows = list(reader)
                fieldnames = reader.fieldnames or []
                encoding = candidate
                break
        except UnicodeDecodeError as exc:
            errors.append(f"{candidate}: {exc}")
    if rows is None:
        raise ValueError("Mine.csv 编码无法读取: " + "; ".join(errors))
    missing = [field for field in REQUIRED_CSV_COLUMNS if field not in fieldnames]
    if missing:
        raise ValueError("Mine.csv 缺失关键字段: " + ", ".join(missing))
    empty_required = {
        field: [row_number for row_number, row in enumerate(rows, start=2) if not str(row.get(field) or "").strip()]
        for field in REQUIRED_CSV_COLUMNS
    }
    empty_required = {field: row_numbers for field, row_numbers in empty_required.items() if row_numbers}
    if empty_required:
        detail = "; ".join(f"{field}: {rows[:10]}" for field, rows in empty_required.items())
        raise ValueError("Mine.csv 必填字段存在空值: " + detail)

    numeric_fields = ("Area", "未治理面积", "Lon", "Lat", "中心经度", "中心纬度")
    invalid_numeric = {}
    for field in numeric_fields:
        if field not in fieldnames:
            continue
        invalid_numeric[field] = [
            row_number
            for row_number, row in enumerate(rows, start=2)
            if str(row.get(field) or "").strip()
            and _to_number(row.get(field)) is None
        ]
    invalid_numeric = {field: row_numbers for field, row_numbers in invalid_numeric.items() if row_numbers}
    if invalid_numeric:
        detail = "; ".join(f"{field}: {rows[:10]}" for field, rows in invalid_numeric.items())
        raise ValueError("Mine.csv 数值字段非法: " + detail)

    duplicate_row_count = len(rows) - len({tuple(row.get(field) for field in fieldnames) for row in rows})
    if duplicate_row_count:
        raise ValueError(f"Mine.csv 存在完全重复行: {duplicate_row_count}")
    duplicate_subject_count = len(rows) - len({str(row.get("主体编号") or "").strip() for row in rows})
    return {
        "field_role": "legacy_subject_code",
        "encoding": encoding,
        "count": len(rows),
        "columns": fieldnames,
        "duplicate_checks": {
            "legacy_subject_code": {
                "count": duplicate_subject_count,
                "note": "主体编号不是权威 TBBH，不参与跨源主键关联",
            },
            "full_row": {"count": duplicate_row_count},
        },
        "empty_checks": {"required_fields": "passed"},
        "numeric_checks": {"checked_fields": [field for field in numeric_fields if field in fieldnames]},
    }


def _validate_manifest_binding(manifest_path, report, paths):
    manifest_file = Path(manifest_path).expanduser().resolve()
    try:
        manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise ValueError(f"资产 manifest 无法读取: {manifest_file}") from exc
    if manifest.get("status") != "ok":
        raise ValueError("资产 manifest 状态不是 ok")
    expected_mapping = ((manifest.get("mapping") or {}).get("tbbh_to_map_fid") or {})
    if expected_mapping != report.get("mapping", {}).get("tbbh_to_map_fid", {}):
        raise ValueError("资产 manifest 的 TBBH/map_fid 映射与当前源文件不一致")
    actual_files = report.get("files", {})
    for path in paths.values():
        matching_entries = [
            entry
            for entry in (manifest.get("files") or {}).values()
            if Path(str(entry.get("path") or "")).name == path.name
        ]
        if len(matching_entries) != 1:
            raise ValueError(f"资产 manifest 缺少唯一文件记录: {path.name}")
        expected = matching_entries[0]
        actual = actual_files.get(
            next(
                name
                for name, item in actual_files.items()
                if Path(str(item.get("path") or "")).name == path.name
            )
        )
        if expected.get("sha256") != actual.get("sha256") or int(expected.get("size", -1)) != int(actual.get("size", -2)):
            raise ValueError(f"资产 manifest 文件哈希或大小不一致: {path.name}")


def validate_jiangxi_assets(
    excel_path=None,
    shp_path=None,
    kmz_path=None,
    csv_path=None,
    expected_count=None,
    manifest_path=None,
):
    defaults = default_asset_paths()
    paths = {
        "excel": Path(excel_path or defaults["excel"]).expanduser().resolve(),
        "shp": Path(shp_path or defaults["shp"]).expanduser().resolve(),
        "kmz": Path(kmz_path or defaults["kmz"]).expanduser().resolve(),
        "csv": Path(csv_path or defaults["csv"]).expanduser().resolve(),
    }
    report = {"status": "failed", "errors": [], "files": {}, "csv": {}}
    try:
        for path in paths.values():
            if not path.is_file():
                raise ValueError(f"缺少文件: {path}")
        excel = _read_excel(paths["excel"])
        shp = _read_shp(paths["shp"], excel["mapping"])
        kmz = _read_kmz(paths["kmz"], excel["mapping"])
        csv_report = _read_csv(paths["csv"])
        set_report = validate_tbbh_sets({
            "excel": excel["tbbh"],
            "shp": shp["tbbh"],
            "kmz": kmz["tbbh"],
        })
        if any(set_report["duplicates"].values()):
            raise ValueError("权威数据存在重复 TBBH")
        if expected_count is not None and any(
            report_count != expected_count
            for report_count in (excel["count"], shp["count"], kmz["count"])
        ):
            raise ValueError(
                "权威数据记录数不符合要求: "
                f"expected={expected_count}, "
                f"Excel={excel['count']}, SHP={shp['count']}, KMZ={kmz['count']}"
            )
        if set_report["intersection"] != sorted(excel["mapping"]):
            raise ValueError("Excel/SHP/KMZ 的 TBBH 集合不一致")
        source_mappings = []
        for source in (shp, kmz):
            if any(item["map_fid"] is None for item in source["records"]):
                raise ValueError("权威数据存在缺失 map_fid")
            source_mappings.append(build_tbbh_map_fid_mapping(source["records"], "tbbh", "map_fid"))
        if source_mappings[0] != source_mappings[1]:
            raise ValueError("SHP/KMZ 的 TBBH 与 map_fid 映射不一致")
        canonical_mapping = source_mappings[0]
        report.update({
            "status": "ok",
            "excel": {"count": excel["count"]},
            "shp": {"count": shp["count"]},
            "kmz": {"count": kmz["count"]},
            "tbbh": {
                **set_report,
                "intersection_count": len(set_report["intersection"]),
            },
            "mapping": {
                "count": len(canonical_mapping),
                "tbbh_to_map_fid": canonical_mapping,
            },
            "csv": csv_report,
        })
        for name, path in paths.items():
            count = {"excel": excel["count"], "shp": shp["count"], "kmz": kmz["count"], "csv": csv_report["count"]}[name]
            report["files"][name] = _file_info(path, count)
        for suffix in (".shx", ".dbf", ".prj"):
            sidecar = paths["shp"].with_suffix(suffix)
            report["files"][sidecar.name] = _file_info(sidecar, shp["count"] if suffix in (".shx", ".dbf") else None)
        if manifest_path:
            _validate_manifest_binding(manifest_path, report, {
                **paths,
                "shx": paths["shp"].with_suffix(".shx"),
                "dbf": paths["shp"].with_suffix(".dbf"),
                "prj": paths["shp"].with_suffix(".prj"),
            })
    except Exception as exc:
        report["status"] = "failed"
        report["errors"].append(str(exc))
    return report


def _build_parser():
    defaults = default_asset_paths()
    parser = argparse.ArgumentParser(description="校验江西 TBBH 权威资产")
    parser.add_argument("--excel", default=str(defaults["excel"]))
    parser.add_argument("--shp", default=str(defaults["shp"]))
    parser.add_argument("--kmz", default=str(defaults["kmz"]))
    parser.add_argument("--csv", default=str(defaults["csv"]))
    parser.add_argument("--expected-count", type=int, default=None)
    parser.add_argument("--manifest", default=None)
    return parser


def main(argv=None):
    args = _build_parser().parse_args(argv)
    explicit_paths = any(flag in (argv or sys.argv[1:]) for flag in ("--excel", "--shp", "--kmz", "--csv"))
    report = validate_jiangxi_assets(
        args.excel,
        args.shp,
        args.kmz,
        args.csv,
        expected_count=args.expected_count if args.expected_count is not None else (None if explicit_paths else 348),
        manifest_path=args.manifest,
    )
    print(f"Excel: {report.get('excel', {}).get('count', 0)}")
    print(f"SHP: {report.get('shp', {}).get('count', 0)}")
    print(f"KMZ: {report.get('kmz', {}).get('count', 0)}")
    print(f"TBBH intersection: {report.get('tbbh', {}).get('intersection_count', 0)}")
    duplicates = report.get("tbbh", {}).get("duplicates", {})
    print(f"duplicate TBBH: {sum(len(values) for values in duplicates.values())}")
    if report.get("errors"):
        for error in report["errors"]:
            print(f"error: {error}")
    print(f"status: {'PASS' if report.get('status') == 'ok' else 'failed'}")
    return 0 if report.get("status") == "ok" else 1


if __name__ == "__main__":
    sys.exit(main())
